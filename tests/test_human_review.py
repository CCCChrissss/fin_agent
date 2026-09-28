import json

import openpyxl
import pytest

from financial_annotation_harness.human_review import (
    GOLD_HEADERS,
    JUDGE_REVIEW_ITEMS,
    RULE_REVIEW_ITEMS,
    read_review_workbook,
    validate_review,
)
from financial_annotation_harness.cli import parser


def expected_record(question_id="Q001"):
    return {
        "question_id": question_id,
        "partition": "development",
        "gold_record_hash": "abc123",
    }


def completed_row(question_id="Q001"):
    return {
        "question_id": question_id,
        "partition": "development",
        "gold_record_hash": "abc123",
        "question_semantics_review": "通過",
        "source_period_unit_review": "通過",
        "evidence_completeness_review": "通過",
        "python_answer_review": "通過",
        "notes": "",
    }


def decisions(items, decision="同意"):
    return [{"item_id": item["item_id"], "decision": decision, "notes": ""} for item in items]


def test_review_is_complete_only_when_all_three_human_gates_are_complete():
    report = validate_review(
        [completed_row()],
        [expected_record()],
        reviewer="研究者",
        rule_reviews=decisions(RULE_REVIEW_ITEMS),
        judge_reviews=decisions(JUDGE_REVIEW_ITEMS),
    )
    assert report["status"] == "COMPLETE"
    assert report["reviewed_questions"] == 1
    assert report["failed_question_ids"] == []
    assert report["blockers"] == []


def test_review_reports_unfinished_cells_and_requires_notes_for_rejection():
    row = completed_row()
    row["evidence_completeness_review"] = "未審核"
    rules = decisions(RULE_REVIEW_ITEMS)
    rules[0]["decision"] = "需調整"
    report = validate_review(
        [row], [expected_record()], reviewer="", rule_reviews=rules,
        judge_reviews=decisions(JUDGE_REVIEW_ITEMS),
    )
    assert report["status"] == "INCOMPLETE"
    assert "reviewer_missing" in report["blockers"]
    assert "question_reviews_incomplete" in report["blockers"]
    assert "rule_review_notes_missing" in report["blockers"]


def test_review_rejects_changed_or_duplicate_gold_identity():
    duplicate = completed_row()
    changed = completed_row()
    changed["gold_record_hash"] = "changed"
    with pytest.raises(ValueError, match="Duplicate question ID"):
        validate_review([duplicate, duplicate], [expected_record()], "研究者", decisions(RULE_REVIEW_ITEMS), decisions(JUDGE_REVIEW_ITEMS))
    with pytest.raises(ValueError, match="Gold record hash mismatch"):
        validate_review([changed], [expected_record()], "研究者", decisions(RULE_REVIEW_ITEMS), decisions(JUDGE_REVIEW_ITEMS))


def test_read_review_workbook_extracts_machine_readable_cells(tmp_path):
    path = tmp_path / "review.xlsx"
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "進度總覽"
    summary["B7"] = "研究者"
    gold = wb.create_sheet("Gold Review")
    gold.append([])
    gold.append([])
    gold.append([])
    gold.append([])
    gold.append(GOLD_HEADERS)
    values = {header: "" for header in GOLD_HEADERS}
    values.update({
        "Question_ID": "Q001", "Partition": "development", "Gold_Record_Hash": "abc123",
        "題意與語意": "通過", "來源資料/單位/期間": "通過", "Evidence完整性": "通過",
        "Python公式與答案": "通過", "審核備註": "",
    })
    gold.append([values[h] for h in GOLD_HEADERS])
    for name, items in (("Rules Review", RULE_REVIEW_ITEMS), ("Judge Review", JUDGE_REVIEW_ITEMS)):
        sheet = wb.create_sheet(name)
        for _ in range(4):
            sheet.append([])
        sheet.append(["Item_ID", "審核項目", "決定", "備註"])
        for item in items:
            sheet.append([item["item_id"], item["label"], "同意", ""])
    wb.save(path)

    parsed = read_review_workbook(path)
    assert parsed["reviewer"] == "研究者"
    assert parsed["gold_reviews"][0]["question_id"] == "Q001"
    assert parsed["rule_reviews"][0]["decision"] == "同意"


def test_cli_exposes_review_validation_and_dev_preflight():
    prepare_args = parser().parse_args(["prepare-review-data", "--output", "review-data.json"])
    validate_args = parser().parse_args(["validate-review", "--input", "review.xlsx", "--output", "report.json"])
    preflight_args = parser().parse_args(["dev-preflight", "--config", "config/experiment.gemma4.v1.json", "--output", "preflight.json"])
    assert prepare_args.command == "prepare-review-data"
    assert validate_args.command == "validate-review"
    assert preflight_args.command == "dev-preflight"
