"""Human-review workbook data and validation. Never mutates the Gold workbook."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from .dataset import inspect_dataset
from .io_utils import digest, file_hash, read_json


GOLD_HEADERS = [
    "Question_ID", "Partition", "Question_Type", "題目", "題意與語意", "來源資料/單位/期間",
    "Evidence完整性", "Python公式與答案", "審核備註", "Overall_Status", "Time", "Concept",
    "Filter", "Logic", "Gold_Answer", "Unit", "Source_Fact_IDs", "來源Facts摘要", "來源儲存格",
    "Golden_Context", "Python_Solution", "Gold_Record_Hash",
]

RULE_REVIEW_ITEMS = [
    {"item_id": "SCHEMA-01", "label": "Schema 與 question_id 必須完整且一致"},
    {"item_id": "FACT-01/02", "label": "Fact 必須存在且來自本題工具回傳"},
    {"item_id": "TIME-01..04", "label": "年度、期間型態與起訖日一致"},
    {"item_id": "UNIT-01..05", "label": "來源單位、答案單位及 context 單位一致"},
    {"item_id": "EVID/NULL", "label": "Evidence 唯一、完整、值正確且不把未揭露視為零"},
    {"item_id": "POT", "label": "Python 可執行、evidence 綁定且答案一致"},
    {"item_id": "TYPE-01/02", "label": "題型與主要運算、Ratio/Growth 公式一致"},
    {"item_id": "COMPLETION", "label": "所有 deterministic validators 通過才可完成"},
]

JUDGE_REVIEW_ITEMS = [
    {"item_id": "JUDGE-INPUT", "label": "Judge 僅看題目、generated annotation、retrieved facts 與 rubric"},
    {"item_id": "JUDGE-FRESH", "label": "Judge 使用獨立 fresh context，不能看到 Generator history 或 hidden reasoning"},
    {"item_id": "JUDGE-NO-GOLD", "label": "Judge 不得看到 Gold answer、Gold label 或 Gold Python result"},
    {"item_id": "JUDGE-RUBRIC", "label": "Judge 檢查語意、evidence 完整性、公式與答案的可支持性"},
    {"item_id": "JUDGE-OUTPUT", "label": "Judge 回傳固定 structured verdict、failure codes 與 feedback"},
    {"item_id": "JUDGE-FAIL-CLOSED", "label": "Judge timeout、malformed 或不完整輸出一律不能通過 completion gate"},
]

_REVIEW_FIELDS = {
    "題意與語意": "question_semantics_review",
    "來源資料/單位/期間": "source_period_unit_review",
    "Evidence完整性": "evidence_completeness_review",
    "Python公式與答案": "python_answer_review",
}


def build_review_payload(root: Path) -> dict:
    """Flatten the immutable Gold source into a review-only derived payload."""
    source = root / "data/financial_qa_gold_dataset_v2.xlsx"
    before = file_hash(source)
    data = inspect_dataset(source)
    split = read_json(root / "artifacts/dataset_split.json")
    partition = {qid: name for name in ("development", "test") for qid in split[name]}
    facts = {row["Fact_ID"]: row for row in data["facts"]}
    questions = {row["Question_ID"]: row for row in data["bank"]}
    rows = []
    for name in ("development", "test"):
        for qid in split[name]:
            gold = next(row for row in data["gold"] if row["Question_ID"] == qid)
            question = questions[qid]
            fact_ids = [value.strip() for value in gold["Source_Fact_IDs"].split(";")]
            selected = [facts[fid] for fid in fact_ids]
            rows.append({
                "question_id": qid,
                "partition": partition[qid],
                "question_type": gold["Question_Type"],
                "question": gold["Chinese_Question"],
                "time": question["Time"],
                "concept": question["Concept"],
                "filter": question["Filter"],
                "logic": question["Logic"],
                "gold_answer": gold["Gold_Answer"],
                "unit": gold["Unit"],
                "source_fact_ids": gold["Source_Fact_IDs"],
                "facts_summary": "\n".join(
                    f'{f["Concept_ZH"]}｜{f["Fiscal_Year"]}｜{f["Value"]}｜{f["Unit"]}｜{f["Period_Start"] or "instant"}~{f["Period_End"]}'
                    for f in selected
                ),
                "source_cells": "; ".join(f'02_Financial_Facts!I{f["_row"]}' for f in selected),
                "golden_context": gold["Golden_Context"],
                "python_solution": gold["Python_Solution"],
                "gold_record_hash": digest({k: v for k, v in gold.items() if k != "_row"}),
            })
    if file_hash(source) != before:
        raise RuntimeError("Original Gold workbook changed while preparing review data")
    return {
        "schema_version": "1.0",
        "purpose": "Human review only; derived from immutable Gold workbook",
        "source_sha256": before,
        "split_seed": split["seed"],
        "rows": rows,
        "rule_review_items": RULE_REVIEW_ITEMS,
        "judge_review_items": JUDGE_REVIEW_ITEMS,
    }


def _sheet_rows(sheet, header_row=5):
    headers = [cell.value for cell in sheet[header_row]]
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in values):
            continue
        yield dict(zip(headers, values))


def read_review_workbook(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        required = {"進度總覽", "Gold Review", "Rules Review", "Judge Review"}
        if not required.issubset(wb.sheetnames):
            raise ValueError(f"Review workbook missing sheets: {sorted(required - set(wb.sheetnames))}")
        reviewer = wb["進度總覽"]["B7"].value or ""
        gold_reviews = []
        for row in _sheet_rows(wb["Gold Review"]):
            gold_reviews.append({
                "question_id": row.get("Question_ID"),
                "partition": row.get("Partition"),
                "gold_record_hash": row.get("Gold_Record_Hash"),
                **{target: row.get(label) for label, target in _REVIEW_FIELDS.items()},
                "notes": row.get("審核備註") or "",
            })
        def decisions(sheet_name):
            return [{"item_id": row.get("Item_ID"), "decision": row.get("決定"), "notes": row.get("備註") or ""}
                    for row in _sheet_rows(wb[sheet_name])]
        return {"reviewer": str(reviewer).strip(), "gold_reviews": gold_reviews,
                "rule_reviews": decisions("Rules Review"), "judge_reviews": decisions("Judge Review")}
    finally:
        wb.close()


def validate_review(gold_reviews: list[dict], expected_records: list[dict], reviewer: str,
                    rule_reviews: list[dict], judge_reviews: list[dict]) -> dict:
    expected = {row["question_id"]: row for row in expected_records}
    seen = set()
    for row in gold_reviews:
        qid = row.get("question_id")
        if qid in seen:
            raise ValueError(f"Duplicate question ID: {qid}")
        seen.add(qid)
        if qid not in expected:
            raise ValueError(f"Unknown question ID: {qid}")
        if row.get("partition") != expected[qid]["partition"]:
            raise ValueError(f"Partition mismatch: {qid}")
        if row.get("gold_record_hash") != expected[qid]["gold_record_hash"]:
            raise ValueError(f"Gold record hash mismatch: {qid}")
    if seen != set(expected):
        raise ValueError(f"Review question set mismatch; missing={sorted(set(expected) - seen)}")

    blockers = []
    if not reviewer.strip():
        blockers.append("reviewer_missing")
    valid_question = {"通過", "不通過"}
    incomplete_questions = [row["question_id"] for row in gold_reviews
                            if any(row.get(field) not in valid_question for field in _REVIEW_FIELDS.values())]
    failed_questions = [row["question_id"] for row in gold_reviews
                        if any(row.get(field) == "不通過" for field in _REVIEW_FIELDS.values())]
    missing_question_notes = [row["question_id"] for row in gold_reviews
                              if row["question_id"] in failed_questions and not str(row.get("notes", "")).strip()]
    if incomplete_questions:
        blockers.append("question_reviews_incomplete")
    if missing_question_notes:
        blockers.append("question_review_notes_missing")

    def validate_decisions(rows, items, prefix):
        by_id = {row.get("item_id"): row for row in rows}
        if set(by_id) != {item["item_id"] for item in items}:
            raise ValueError(f"{prefix} review item set mismatch")
        if any(row.get("decision") not in {"同意", "需調整"} for row in rows):
            blockers.append(f"{prefix}_reviews_incomplete")
        if any(row.get("decision") == "需調整" and not str(row.get("notes", "")).strip() for row in rows):
            blockers.append(f"{prefix}_review_notes_missing")
        return [row["item_id"] for row in rows if row.get("decision") == "需調整"]

    failed_rules = validate_decisions(rule_reviews, RULE_REVIEW_ITEMS, "rule")
    failed_judge = validate_decisions(judge_reviews, JUDGE_REVIEW_ITEMS, "judge")
    needs_changes = bool(failed_questions or failed_rules or failed_judge)
    status = "INCOMPLETE" if blockers else "NEEDS_CHANGES" if needs_changes else "COMPLETE"
    return {
        "schema_version": "1.0", "status": status, "reviewer": reviewer.strip(),
        "question_count": len(expected), "reviewed_questions": len(expected) - len(incomplete_questions),
        "failed_question_ids": failed_questions, "incomplete_question_ids": incomplete_questions,
        "failed_rule_ids": failed_rules, "failed_judge_item_ids": failed_judge,
        "blockers": blockers,
    }


def validate_review_workbook(path: Path, root: Path) -> dict:
    parsed = read_review_workbook(path)
    payload = build_review_payload(root)
    expected = [{k: row[k] for k in ("question_id", "partition", "gold_record_hash")} for row in payload["rows"]]
    report = validate_review(parsed["gold_reviews"], expected, parsed["reviewer"],
                             parsed["rule_reviews"], parsed["judge_reviews"])
    return {**report, "source_sha256": payload["source_sha256"], "split_seed": payload["split_seed"]}

